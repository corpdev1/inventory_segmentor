"""Excel I/O for the inventory segmenter."""

from __future__ import annotations

import pandas as pd
import re

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BUCKETS: list[tuple[int, str]] = [
    (1, "Product & Engineering"),
    (2, "Customer & Sales"),
    (3, "Strategy & Planning"),
    (4, "Financial & Legal"),
    (5, "Operations & HR"),
    (6, "Marketing"),
    (7, "Meeting Notes & Internal Comms"),
]

DESCRIPTION_COL_CANDIDATES = [
    "description",
    "what it includes",
    "item",
    "details",
    "notes",
    "content",
    "data",
]
SOURCE_COL_CANDIDATES = [
    "source",
    "where to find it",
    "system",
    "tool",
    "platform",
    "location",
]


def _normalize(value: object) -> str:
    return str(value).strip().lower()


def detect_column(df: pd.DataFrame, candidates: list[str]) -> str | None:
    """Return the first column matching any candidate (exact, then substring)."""
    cols_normalized = {_normalize(c): c for c in df.columns}
    for cand in candidates:
        if cand in cols_normalized:
            return cols_normalized[cand]
    for cand in candidates:
        for norm, original in cols_normalized.items():
            if cand in norm:
                return original
    return None


def load_inventory(
    path: str,
    sheet_name: str | None = None,
    description_column: str | None = None,
    source_column: str | None = None,
) -> tuple[pd.DataFrame, str, str | None]:
    """Load the inventory; return (df_with_row_id, description_col, source_col)."""
    if path.lower().endswith(".csv"):
        df = pd.read_csv(path)
    else:
        df = pd.read_excel(path, sheet_name=sheet_name or 0)

    if description_column is None:
        description_column = detect_column(df, DESCRIPTION_COL_CANDIDATES)
    if description_column is None:
        # last resort: first text column
        for col in df.columns:
            if df[col].dtype == object:
                description_column = col
                break
    if description_column is None:
        raise ValueError(
            "Could not identify a description column. "
            "Pass description_column explicitly when calling the tool."
        )

    if source_column is None:
        source_column = detect_column(df, SOURCE_COL_CANDIDATES)

    if "row_id" not in df.columns:
        df.insert(0, "row_id", range(1, len(df) + 1))

    return df, description_column, source_column


def _write_sheet(ws, frame: pd.DataFrame) -> None:
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="305496")
    header_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    body_font = Font(name="Arial")
    body_align = Alignment(vertical="top", wrap_text=True)

    cols = list(frame.columns)
    ws.append(cols)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    def _clean_cell_value(v):
        # openpyxl will throw IllegalCharacterError for certain control chars
        # (common when extracting text from PDFs). Strip them.
        if isinstance(v, str):
            return ILLEGAL_CHARACTERS_RE.sub("", v)
        return v

    for record in frame.itertuples(index=False):
        ws.append([_clean_cell_value(v) for v in list(record)])

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = body_align

    for i, col in enumerate(cols, start=1):
        sample = frame[col].astype(str).tolist()[:200] if len(frame) else []
        max_len = (
            max([len(str(col))] + [len(str(v)) for v in sample])
            if sample
            else len(str(col))
        )
        ws.column_dimensions[get_column_letter(i)].width = min(max(14, max_len + 2), 60)
    ws.freeze_panes = "A2"


def write_segmented_workbook(df: pd.DataFrame, output_path: str) -> None:
    """Write Master + 7 bucket sheets + Summary to output_path."""
    wb = Workbook()
    wb.remove(wb.active)

    ws_master = wb.create_sheet("Master")
    _write_sheet(ws_master, df)

    for num, name in BUCKETS:
        bucket_df = df[df["bucket_number"] == num].copy()
        sheet_title = f"{num}. {name}"[:31]  # Excel's 31-char sheet name limit
        ws = wb.create_sheet(sheet_title)
        _write_sheet(ws, bucket_df)

    # Summary
    ws_summary = wb.create_sheet("Summary")
    summary_cols = ["Bucket #", "Bucket Name", "Items", "High Confidence", "Low Confidence"]
    ws_summary.append(summary_cols)
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="305496")
    for cell in ws_summary[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")

    for num, name in BUCKETS:
        bucket_df = df[df["bucket_number"] == num]
        high = int((bucket_df.get("confidence", pd.Series([], dtype=str)) == "high").sum())
        low = int((bucket_df.get("confidence", pd.Series([], dtype=str)) == "low").sum())
        ws_summary.append([num, name, len(bucket_df), high, low])

    total_row = len(BUCKETS) + 2
    ws_summary.cell(row=total_row, column=1, value="Total").font = Font(name="Arial", bold=True)
    ws_summary.cell(row=total_row, column=3, value=f"=SUM(C2:C{total_row - 1})")
    ws_summary.cell(row=total_row, column=4, value=f"=SUM(D2:D{total_row - 1})")
    ws_summary.cell(row=total_row, column=5, value=f"=SUM(E2:E{total_row - 1})")

    body_font = Font(name="Arial")
    for row in ws_summary.iter_rows(min_row=2):
        for cell in row:
            if cell.font.bold:
                continue
            cell.font = body_font
    for i in range(1, len(summary_cols) + 1):
        ws_summary.column_dimensions[get_column_letter(i)].width = 24
    ws_summary.freeze_panes = "A2"

    wb.save(output_path)


def write_company_inventory_workbook(
    *,
    inventory_df: pd.DataFrame,
    evidence_df: pd.DataFrame | None,
    output_path: str,
) -> None:
    """Write a clean multi-tab workbook for dump-derived outputs.

    Sheets:
      - Inventory: 7-row desired format
      - Overview: quick counts by bucket/confidence/source/ext (if evidence provided)
      - Sample Files: first N evidence rows (easy browsing)
      - Evidence: full evidence table (if provided)
      - One sheet per bucket: "<bucket name>" with evidence rows for that bucket
    """
    wb = Workbook()
    wb.remove(wb.active)

    ws_inventory = wb.create_sheet("Inventory")
    _write_sheet(ws_inventory, inventory_df)

    if evidence_df is not None:
        # Build a clean, user-facing evidence view (like the screenshot).
        edf = evidence_df.copy()

        def _size_human(n: object) -> str:
            try:
                x = float(n)
            except Exception:
                return ""
            units = ["B", "KB", "MB", "GB", "TB"]
            i = 0
            while x >= 1024 and i < len(units) - 1:
                x /= 1024
                i += 1
            if i == 0:
                return f"{int(x)} {units[i]}"
            return f"{x:.1f} {units[i]}"

        display = pd.DataFrame(
            {
                "Item Name": edf.get("filename", ""),
                "File Type": edf.get("extension", ""),
                "Content Summary": edf.get("content_summary", ""),
                "PII Flag": edf.get("pii_flag", ""),
                "Size": edf.get("size_bytes", "").map(_size_human)
                if "size_bytes" in edf.columns
                else "",
                "Source": edf.get("source_guess", ""),
                "Path / Subsection": edf.get("path", ""),
                # Keep bucket info for filtering/tabs + auditability.
                "bucket_number": edf.get("bucket_number", ""),
                "bucket_name": edf.get("bucket_name", ""),
                "confidence": edf.get("confidence", ""),
                "rationale": edf.get("rationale", ""),
            }
        )

        # Overview
        ws_overview = wb.create_sheet("Overview")
        overview_rows: list[dict[str, object]] = []

        # Bucket summary
        if "bucket_number" in evidence_df.columns:
            bucket_counts = (
                evidence_df["bucket_number"]
                .value_counts(dropna=False)
                .sort_index()
                .to_dict()
            )
            for k, v in bucket_counts.items():
                overview_rows.append({"Metric": "bucket_count", "Key": str(k), "Value": int(v)})

        # Confidence summary
        if "confidence" in evidence_df.columns:
            conf_counts = evidence_df["confidence"].value_counts(dropna=False).to_dict()
            for k, v in conf_counts.items():
                overview_rows.append({"Metric": "confidence_count", "Key": str(k), "Value": int(v)})

        # Source guesses
        if "source_guess" in evidence_df.columns:
            src_counts = (
                evidence_df["source_guess"]
                .fillna("")
                .astype(str)
                .replace({"nan": ""})
                .value_counts()
                .head(40)
                .to_dict()
            )
            for k, v in src_counts.items():
                key = k if k else "(blank)"
                overview_rows.append({"Metric": "source_guess_top", "Key": key, "Value": int(v)})

        # File extensions
        if "extension" in evidence_df.columns:
            ext_counts = (
                evidence_df["extension"]
                .fillna("")
                .astype(str)
                .replace({"nan": ""})
                .str.lower()
                .value_counts()
                .head(40)
                .to_dict()
            )
            for k, v in ext_counts.items():
                key = k if k else "(blank)"
                overview_rows.append({"Metric": "extension_top", "Key": key, "Value": int(v)})

        if overview_rows:
            _write_sheet(ws_overview, pd.DataFrame(overview_rows))
        else:
            _write_sheet(ws_overview, pd.DataFrame([{"Metric": "info", "Key": "", "Value": "No evidence provided"}]))

        # Sample Files (first N rows, stable + human-friendly)
        ws_sample = wb.create_sheet("Sample Files")
        sample_n = 250 if len(display) >= 250 else len(display)
        _write_sheet(ws_sample, display.head(sample_n))

        # Full Evidence
        ws_evidence = wb.create_sheet("Evidence")
        _write_sheet(ws_evidence, display)

        # One tab per bucket with that bucket's evidence (clean browsing)
        if "bucket_number" in display.columns and "bucket_name" in display.columns:
            for num, name in BUCKETS:
                bdf = display[display["bucket_number"] == num].copy()
                if len(bdf) == 0:
                    continue
                # Prefer the plain bucket name as sheet title (what users expect)
                sheet_title = str(name)[:31]
                # Ensure uniqueness if names collide after truncation
                if sheet_title in wb.sheetnames:
                    sheet_title = f"{num}. {name}"[:31]
                ws = wb.create_sheet(sheet_title)
                _write_sheet(ws, bdf)

    wb.save(output_path)
